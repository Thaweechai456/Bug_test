import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_font_run(run, font_name="TH Sarabun PSK", size_pt=16, bold=None, color_rgb=None):
    if bold is not None:
        run.bold = bold
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if color_rgb is not None:
        run.font.color.rgb = color_rgb
        
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
        
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    
    lang = rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rPr.append(lang)
    lang.set(qn('w:val'), 'th-TH')
    lang.set(qn('w:bidi'), 'th-TH')
    lang.set(qn('w:eastAsia'), 'th-TH')

def enforce_doc_fonts(doc, default_font="TH Sarabun PSK"):
    for p in doc.paragraphs:
        for r in p.runs:
            set_font_run(r, font_name=default_font)
            
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_font_run(r, font_name=default_font)

def clean_and_build_workshop1():
    doc = docx.Document('Workshop_Developer_to_Tester_completed.docx')
    
    # Table 0: Student info
    t0 = doc.tables[0]
    t0.rows[1].cells[0].text = ''
    r1 = t0.rows[1].cells[0].paragraphs[0].add_run('66040233102')
    set_font_run(r1, bold=True, size_pt=16)
    
    t0.rows[1].cells[1].text = ''
    r2 = t0.rows[1].cells[1].paragraphs[0].add_run('นายทวีชัย นามฮาด')
    set_font_run(r2, bold=True, size_pt=16)
    
    cell_img = t0.rows[1].cells[2]
    cell_img.text = ''
    p_img = cell_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists('screenshots/01_website_header.png'):
        p_img.add_run().add_picture('screenshots/01_website_header.png', width=Inches(2.5))

    # Remove any dangling paragraphs with only 'ภาพหลักฐานการ Re-test' at the bottom
    for p in doc.paragraphs:
        if p.text.strip() == 'ภาพหลักฐานการ Re-test หลังแก้ไขโค้ด (PASS ทุกเงื่อนไข):':
            p.text = ''

    # Workshop 1 Checklist clean up
    checklist_items = [
        'ฉันอธิบายการทำงานของส่วนที่ได้รับมอบหมายได้',
        'ฉันสร้าง Test Data และ Expected Result ได้',
        'ฉันทดลองระบบจริงและบันทึก Actual Result',
        'ฉันพบ Bug และมีหลักฐาน',
        'ฉันเขียน Bug Report ให้ผู้อื่นทำตามได้',
        'ฉัน Re-test และสรุป PASS / FAIL ได้'
    ]
    
    chk_idx = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if any(item[:10] in txt for item in checklist_items) or ('ฉัน' in txt and ('PASS' in txt or 'Re-test' in txt or 'Bug' in txt or 'Test Data' in txt or 'Actual' in txt or 'อธิบาย' in txt)):
            if chk_idx < len(checklist_items):
                p.text = ''
                r_chk = p.add_run('[✓] ' + checklist_items[chk_idx])
                set_font_run(r_chk, bold=True, size_pt=16, color_rgb=RGBColor(20, 83, 45))
                chk_idx += 1

    enforce_doc_fonts(doc)
    doc.save('Workshop_Developer_to_Tester_completed.docx')
    print('Workshop 1 cleaned and saved!')

def clean_and_build_workshop2():
    doc = docx.Document('Workshop2_DataFlowDetection_completed.docx')
    
    # Fix Reflection Question 2 (remove duplicate text)
    for i, p in enumerate(doc.paragraphs):
        if '2. หลังทำกิจกรรมนี้ ฉันเห็นความสัมพันธ์' in p.text:
            if i + 1 < len(doc.paragraphs):
                p_ans2 = doc.paragraphs[i+1]
                p_ans2.text = ''
                r_ans2 = p_ans2.add_run(
                    'Source Code บอก Def/Use และ branch ส่วน Test Case เลือก Input เพื่อให้เดินผ่าน Path ที่ต้องการพิสูจน์\n'
                    'phone เป็นตัวอย่างที่ทำให้เห็นความเชื่อมโยงจาก input ไปสู่ regex และผล validation\n'
                    'ดังนั้น Test Case สามารถสร้างจากโครงสร้างของโปรแกรมได้'
                )
                set_font_run(r_ans2, size_pt=16)
            # check if i+2 is duplicate
            if i + 2 < len(doc.paragraphs) and 'phone เป็นตัวอย่าง' in doc.paragraphs[i+2].text:
                doc.paragraphs[i+2].text = ''
            break

    enforce_doc_fonts(doc)
    doc.save('Workshop2_DataFlowDetection_completed.docx')
    print('Workshop 2 cleaned and saved!')

if __name__ == '__main__':
    clean_and_build_workshop1()
    clean_and_build_workshop2()
