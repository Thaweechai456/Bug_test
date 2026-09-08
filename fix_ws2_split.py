import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_font_run(run, font_name="TH Sarabun PSK", size_pt=16, bold=None, color_rgb=None):
    if bold is not None:
        run.bold = bold
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if color_rgb is not None:
        run.font.color.rgb = color_rgb
        
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
        
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    
    lang = rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rPr.append(lang)
    lang.set(qn('w:val'), 'th-TH')
    lang.set(qn('w:bidi'), 'th-TH')
    lang.set(qn('w:eastAsia'), 'th-TH')

def enforce_doc_fonts(doc, default_font="TH Sarabun PSK"):
    for p in doc.paragraphs:
        for r in p.runs:
            set_font_run(r, font_name=default_font)
            
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_font_run(r, font_name=default_font)

def clean_and_build_workshop2():
    doc = docx.Document('Workshop2_DataFlowDetection_completed.docx')
    
    for i, p in enumerate(doc.paragraphs):
        if '2. หลังทำกิจกรรมนี้ ฉันเห็นความสัมพันธ์' in p.text:
            doc.paragraphs[i+1].text = 'Source Code บอก Def/Use และ branch ส่วน Test Case เลือก Input เพื่อให้เดินผ่าน Path ที่ต้องการพิสูจน์'
            doc.paragraphs[i+2].text = 'phone เป็นตัวอย่างที่ทำให้เห็นความเชื่อมโยงจาก input ไปสู่ regex และผล validation'
            doc.paragraphs[i+3].text = 'ดังนั้น Test Case สามารถสร้างจากโครงสร้างของโปรแกรมได้'
            break

    enforce_doc_fonts(doc)
    doc.save('Workshop2_DataFlowDetection_completed.docx')
    print('Workshop 2 perfected!')

if __name__ == '__main__':
    clean_and_build_workshop2()
