import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def update_workshop1():
    doc = docx.Document('Workshop_Developer_to_Tester_completed.docx')
    
    # 1. Table 0: Fix Student ID and Add Profile Image
    t0 = doc.tables[0]
    t0.rows[1].cells[0].text = '66040233102'
    t0.rows[1].cells[1].text = 'นายทวีชัย นามฮาด'
    cell_img = t0.rows[1].cells[2]
    cell_img.text = ''
    p_img = cell_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists('screenshots/01_website_header.png'):
        p_img.add_run().add_picture('screenshots/01_website_header.png', width=Inches(2.5))
        
    # 2. Insert Website Overview at P11/P12
    for i, p in enumerate(doc.paragraphs):
        if '1.1 นำภาพหน้าเว็บไซต์ของนักศึกษาใส่ไว้ที่นี่' in p.text:
            p.text = '1.1 ภาพหน้าเว็บไซต์ของผู้พัฒนา (Developer Website Overview):'
            # Check next paragraph or insert picture
            if i + 1 < len(doc.paragraphs) and doc.paragraphs[i+1].text.strip() == '':
                p_target = doc.paragraphs[i+1]
            else:
                p_target = p.insert_paragraph_before()
            p_target.text = ''
            p_target.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists('screenshots/01_website_overview.png'):
                p_target.add_run().add_picture('screenshots/01_website_overview.png', width=Inches(5.5))
            break

    # 3. Evidence in Part 3 (P47/P48)
    for i, p in enumerate(doc.paragraphs):
        if '[แทรก Screenshot ของข้อความ Bug Spoiler' in p.text or 'Evidence: หน้า https://' in p.text:
            p.text = 'Evidence: ตรวจสอบจากหน้าเว็บส่วน Bug Spoilers และทดสอบกรอกเบอร์โทรศัพท์ 10 หลักพบ Error ดังภาพ:'
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_pic1 = p.insert_paragraph_before()
            p_pic1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists('screenshots/02_bug_spoilers.png'):
                p_pic1.add_run().add_picture('screenshots/02_bug_spoilers.png', width=Inches(5.2))
            p_pic2 = p.insert_paragraph_before()
            p_pic2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists('screenshots/03_bug1_phone_10digits_fail.png'):
                p_pic2.add_run().add_picture('screenshots/03_bug1_phone_10digits_fail.png', width=Inches(5.2))
            break

    # 4. Table 2 (Bug Report - BUG-01) Evidence
    t2 = doc.tables[2]
    cell_ev = t2.rows[5].cells[1]
    cell_ev.text = 'หลักฐานข้อผิดพลาด (Regex ตรวจสอบ 9 หลัก แทนที่จะเป็น 10 หลัก):\n'
    p_ev = cell_ev.paragraphs[0]
    if os.path.exists('screenshots/03_bug1_phone_10digits_fail.png'):
        p_ev.add_run().add_picture('screenshots/03_bug1_phone_10digits_fail.png', width=Inches(4.5))

    # 5. Part 5: Re-test
    for i, p in enumerate(doc.paragraphs):
        if '5.1 หลังจาก Developer แก้ไข' in p.text:
            doc.paragraphs[i+1].text = 'ทำการ Re-test ทดสอบกรณีเดิมซ้ำทั้งหมด 100% หลังแก้ไข Source Code ใน script.js ทั้งจุดเบอร์โทรศัพท์และจุดกดยอมรับเงื่อนไข'
        elif '5.2 ผลการทดสอบเป็น PASS หรือ FAIL?' in p.text:
            doc.paragraphs[i+1].text = 'ผลการ Re-test: PASS ทั้งหมด (ทุกกรณีทำงานถูกต้องสมบูรณ์ตามเกณฑ์มาตรฐาน)'
        elif '5.3 เพราะเหตุใดจึงสรุปผลเช่นนั้น?' in p.text:
            doc.paragraphs[i+1].text = 'สรุปผลจากการทดสอบจริงหลังแก้ไขโค้ด:\n' \
                '1) จุดที่ 1 (validatePhone): แก้ไข Regex เป็น /^0[0-9]{9}$/ ส่งผลให้เบอร์ 10 หลัก (0812345678) ผ่านการตรวจสอบ (PASS) และเมื่อกรอก 9 หลัก (081234567) ระบบจะแจ้งเตือนว่าผิด (PASS)\n' \
                '2) จุดที่ 2 (validateTerms): ลบตัวแปร Bypass และตรวจ !termsCheckbox.checked จริง ส่งผลให้หากไม่ติ๊กถูก ระบบจะแสดงข้อความเตือนและไม่อนุญาตให้ Submit (PASS) และเมื่อติ๊กถูกจะลงทะเบียนสำเร็จ (PASS)'
        elif 'แนบ Screenshot ของผลหลังแก้ไข' in p.text:
            p.text = '📷 ภาพหลักฐานการ Re-test หลังแก้ไขโค้ด (PASS ทุกเงื่อนไข):'
            p1 = p.insert_paragraph_before()
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.add_run('1. Re-test เบอร์ 9 หลัก -> ระบบตรวจพบและแจ้งเตือนถูกต้อง (PASS)\n')
            if os.path.exists('screenshots/07_retest_phone_9digits_rejected.png'):
                p1.add_run().add_picture('screenshots/07_retest_phone_9digits_rejected.png', width=Inches(4.8))
                
            p2 = p.insert_paragraph_before()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.add_run('\n2. Re-test ไม่ติ๊กยอมรับเงื่อนไข -> ระบบปฏิเสธการส่งฟอร์ม (PASS)\n')
            if os.path.exists('screenshots/09_retest_terms_unchecked_rejected.png'):
                p2.add_run().add_picture('screenshots/09_retest_terms_unchecked_rejected.png', width=Inches(4.8))

            p3 = p.insert_paragraph_before()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p3.add_run('\n3. Re-test กรอกข้อมูลถูกต้องครบถ้วน 10 หลัก + ติ๊กยอมรับเงื่อนไข -> ส่งสำเร็จ (PASS)\n')
            if os.path.exists('screenshots/10_retest_all_valid_submit_success.png'):
                p3.add_run().add_picture('screenshots/10_retest_all_valid_submit_success.png', width=Inches(4.8))

    # 6. Checklist checkmarks
    for p in doc.paragraphs:
        if '☐ ฉัน' in p.text:
            p.text = p.text.replace('☐', '☑')

    doc.save('Workshop_Developer_to_Tester_completed.docx')
    print('Workshop 1 updated successfully!')

def update_workshop2():
    doc = docx.Document('Workshop2_DataFlowDetection_completed.docx')
    
    # 1. Table 0: Student Info & Header image
    t0 = doc.tables[0]
    t0.rows[1].cells[0].text = '66040233102'
    t0.rows[1].cells[1].text = 'นายทวีชัย นามฮาด'
    cell_img = t0.rows[1].cells[2]
    cell_img.text = ''
    p_img = cell_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists('screenshots/01_website_header.png'):
        p_img.add_run().add_picture('screenshots/01_website_header.png', width=Inches(2.5))
        
    # 2. Table 2: Website Overview
    t2 = doc.tables[2]
    cell_t2 = t2.rows[0].cells[0]
    cell_t2.text = ''
    p_t2 = cell_t2.paragraphs[0]
    p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists('screenshots/01_website_overview.png'):
        p_t2.add_run().add_picture('screenshots/01_website_overview.png', width=Inches(5.5))
        
    # 3. Mission 2 Step 1: Bug Spoilers Screenshot
    for p in doc.paragraphs:
        if '[แทรก Screenshot ของ Source/ Bug Spoiler ที่นี่]' in p.text:
            p.text = 'หลักฐานจาก Bug Spoilers บนเว็บไซต์:'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists('screenshots/02_bug_spoilers.png'):
                p.add_run('\n').add_picture('screenshots/02_bug_spoilers.png', width=Inches(5.0))
        elif '[แทรก DFG แบบลูกศร/ Screenshot Source Code ที่นี่]' in p.text:
            p.text = 'โครงสร้าง Data Flow Graph (DFG) ของตัวแปร phone:'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists('screenshots/dfg_diagram.png'):
                p.add_run('\n').add_picture('screenshots/dfg_diagram.png', width=Inches(5.2))

    # 4. Table 13: Test Results Evidence
    t13 = doc.tables[13]
    # TC-01
    c1 = t13.rows[1].cells[3]
    c1.text = ''
    if os.path.exists('screenshots/04_bug1_phone_9digits_pass_modal.png'):
        c1.paragraphs[0].add_run().add_picture('screenshots/04_bug1_phone_9digits_pass_modal.png', width=Inches(2.2))
    # TC-02
    c2 = t13.rows[2].cells[3]
    c2.text = ''
    if os.path.exists('screenshots/03_bug1_phone_10digits_fail.png'):
        c2.paragraphs[0].add_run().add_picture('screenshots/03_bug1_phone_10digits_fail.png', width=Inches(2.2))
    # TC-03
    t13.rows[3].cells[1].text = 'ระบบแสดง phone error แจ้งเตือนรูปแบบไม่ถูกต้อง'
    t13.rows[3].cells[2].text = 'PASS'
    c3 = t13.rows[3].cells[3]
    c3.text = ''
    if os.path.exists('screenshots/06_phone_invalid_letters_error.png'):
        c3.paragraphs[0].add_run().add_picture('screenshots/06_phone_invalid_letters_error.png', width=Inches(2.2))
    # TC-04
    c4 = t13.rows[4].cells[3]
    c4.text = ''
    if os.path.exists('screenshots/02_bug_spoilers.png'):
        c4.paragraphs[0].add_run().add_picture('screenshots/02_bug_spoilers.png', width=Inches(2.2))

    # 5. Table 15: Bug Report Evidence
    t15 = doc.tables[15]
    c15 = t15.rows[8].cells[1]
    c15.text = 'หลักฐานจาก Bug Spoilers และการทดสอบ:\n'
    if os.path.exists('screenshots/03_bug1_phone_10digits_fail.png'):
        c15.paragraphs[0].add_run().add_picture('screenshots/03_bug1_phone_10digits_fail.png', width=Inches(4.2))

    # 6. Table 16: Re-test Results & Evidence
    t16 = doc.tables[16]
    for r in range(1, 4):
        t16.rows[r].cells[1].text = ''
        t16.rows[r].cells[2].text = ''
    t16.rows[1].cells[1].text = 'ใช้ Test Data เดิม: ทดสอบเบอร์ 9 หลัก, 10 หลัก และกรณีไม่ติ๊กยอมรับเงื่อนไข'
    t16.rows[1].cells[2].text = 'ใช้ Test Data เดิม: ทดสอบเบอร์ 9 หลัก, 10 หลัก และกรณีไม่ติ๊กยอมรับเงื่อนไข'
    
    t16.rows[2].cells[1].text = 'หลังแก้ไขโค้ด: เบอร์ 10 หลักผ่านการตรวจสอบ, เบอร์ 9 หลักขึ้น Error ถูกต้อง, และไม่ติ๊ก Terms จะส่งฟอร์มไม่ได้'
    t16.rows[2].cells[2].text = 'หลังแก้ไขโค้ด: เบอร์ 10 หลักผ่านการตรวจสอบ, เบอร์ 9 หลักขึ้น Error ถูกต้อง, และไม่ติ๊ก Terms จะส่งฟอร์มไม่ได้'
    
    t16.rows[3].cells[1].text = 'PASS (ผ่านการทดสอบ 100%)'
    t16.rows[3].cells[2].text = 'PASS (ผ่านการทดสอบ 100%)'
    
    c16 = t16.rows[4].cells[1]
    c16.text = 'ภาพหลักฐาน Re-test สำเร็จ:\n'
    if os.path.exists('screenshots/10_retest_all_valid_submit_success.png'):
        c16.paragraphs[0].add_run().add_picture('screenshots/10_retest_all_valid_submit_success.png', width=Inches(4.0))
    t16.rows[4].cells[2].text = ''
    if os.path.exists('screenshots/07_retest_phone_9digits_rejected.png'):
        t16.rows[4].cells[2].paragraphs[0].add_run('หลักฐานปฏิเสธเบอร์ 9 หลัก:\n').add_picture('screenshots/07_retest_phone_9digits_rejected.png', width=Inches(4.0))

    # 7. Checklist checkmarks
    for p in doc.paragraphs:
        if '[ ]' in p.text:
            p.text = p.text.replace('[ ]', '[✓]')

    doc.save('Workshop2_DataFlowDetection_completed.docx')
    print('Workshop 2 updated successfully!')

if __name__ == '__main__':
    update_workshop1()
    update_workshop2()
