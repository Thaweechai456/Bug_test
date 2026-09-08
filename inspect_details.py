import docx

def dump_detailed(doc_path, out_path):
    doc = docx.Document(doc_path)
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(f'=== PARAGRAPHS FOR {doc_path} ===\n')
        for i, p in enumerate(doc.paragraphs):
            f.write(f'P[{i}]: {p.text}\n')
        
        f.write(f'\n=== TABLES FOR {doc_path} ===\n')
        for t_idx, t in enumerate(doc.tables):
            f.write(f'\nTable {t_idx} ({len(t.rows)}x{len(t.columns)}):\n')
            for r_idx, row in enumerate(t.rows):
                for c_idx, cell in enumerate(row.cells):
                    cell_p = [p.text for p in cell.paragraphs]
                    f.write(f'  Cell[{r_idx},{c_idx}]: {" || ".join(cell_p)}\n')

dump_detailed('Workshop_Developer_to_Tester_completed.docx', 'doc1_detailed.txt')
dump_detailed('Workshop2_DataFlowDetection_completed.docx', 'doc2_detailed.txt')
print('Detailed dumps complete')
